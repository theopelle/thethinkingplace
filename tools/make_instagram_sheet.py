"""The Thinking Place: the Instagram calendar spreadsheet.

    python3 tools/make_instagram_sheet.py

Reads exports/instagram/posts.json (written by tools/make_instagram.js) and writes
exports/instagram/instagram-calendar.xlsx: a "How to use" sheet and one row per post with
its caption and alt text. The "Scheduled" column is the user's: running this again keeps
what they typed there.
"""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "exports" / "instagram"
XLSX = OUT / "instagram-calendar.xlsx"
FONT = "Arial"
TEAL, STONE, PAPER, YELLOW = "4E7378", "E3DFD7", "FBFAF7", "FFF4C2"
FORMAT = {"quote": "Quote (1 image)", "carousel": "Carousel (4 slides)", "question": "Question card (1 image)"}

COLUMNS = [  # header, width
    ("Date", 12), ("Weekday", 11), ("Month", 9), ("Format", 21), ("Mode", 7), ("Thought #", 9),
    ("Author", 24), ("Path", 24), ("Occasion", 18), ("Images folder", 24), ("Caption", 70),
    ("Alt text, slide 1", 40), ("Alt text, slide 2", 40), ("Alt text, slide 3", 40),
    ("Alt text, slide 4", 40), ("Scheduled", 11),
]

HOW_TO = [
    ("The Thinking Place on Instagram", None),
    ("One post on Mondays, Wednesdays and Fridays for a year, each with the thought the app shows that day, "
     "so someone who follows the link finds the same thought in the app.", None),
    ("", None),
    ("The rhythm", "h"),
    ("Monday: the quote as one image. Wednesday: a carousel (quote, about this thought, the question, end card). "
     "Friday: the question as one image. Pinned dates (the longest night, the New Year, the equinoxes, blossom "
     "time, midsummer) are always a carousel. Posts alternate light and dark, so the profile grid is a checkerboard.", None),
    ("A revisit day (from April 2027) posts the revisited thought only if it has not been on Instagram yet; "
     "otherwise that day is left out. Thoughts that mention God are left out too.", None),
    ("", None),
    ("Scheduling a month at a time (Meta Business Suite, free)", "h"),
    ("1. On a computer, open business.facebook.com → Planner (or Content) → Create post, with Instagram selected.", None),
    ("2. Filter the Posts sheet by Month. For each row, add the PNGs from its images folder in order "
     "(slide-1, slide-2, …). Two or more images become a carousel.", None),
    ("3. Paste the Caption cell. Under alt text (Advanced settings), paste the matching Alt text cells.", None),
    ("4. Schedule for the row's date. Morning, about 7:30, is a good default, the hour the app's reminders go out.", None),
    ("5. Write \"yes\" in Scheduled. The count below goes up.", None),
    ("The Instagram app can schedule the same way (Advanced settings → Schedule) if you prefer the phone.", None),
    ("", None),
    ("Before each batch", "h"),
    ("Glance at the month in index.html (same folder): every post, full size, with its caption.", None),
    ("When the Android app is live on Google Play, change APP_LINE in tools/make_instagram.js and run the two "
     "commands below again: every caption updates, and the Scheduled column is kept.", None),
    ("node tools/make_instagram.js 2026-10-26 52", "code"),
    ("python3 tools/make_instagram_sheet.py", "code"),
    ("", None),
    ("Scheduled so far", "h"),
]


def scheduled_before():
    if not XLSX.exists():
        return {}
    ws = load_workbook(XLSX)["Posts"]
    head = [c.value for c in ws[1]]
    i_date, i_id, i_s = head.index("Date"), head.index("Thought #"), head.index("Scheduled")
    return {(str(r[i_date]), r[i_id]): r[i_s] for r in ws.iter_rows(min_row=2, values_only=True) if r[i_s]}


def main():
    posts = json.loads((OUT / "posts.json").read_text())
    kept = scheduled_before()
    wb = Workbook()
    how = wb.active
    how.title = "How to use"
    how.column_dimensions["A"].width = 110
    how.sheet_view.showGridLines = False
    for n, (text, kind) in enumerate(HOW_TO, 1):
        c = how.cell(n, 1, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=FONT, size=11)
        if n == 1:
            c.font = Font(name=FONT, size=16, bold=True, color=TEAL)
        elif kind == "h":
            c.font = Font(name=FONT, size=12, bold=True, color=TEAL)
        elif kind == "code":
            c.font = Font(name="Courier New", size=10)
    n = len(HOW_TO) + 1
    how.cell(n, 1, f'=COUNTIF(Posts!P:P,"yes")&" of "&COUNTA(Posts!A:A)-1&" posts"').font = Font(name=FONT, size=11)

    ws = wb.create_sheet("Posts")
    for i, (name, width) in enumerate(COLUMNS, 1):
        c = ws.cell(1, i, name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TEAL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    for r, p in enumerate(posts, 2):
        alts = p["alt"] + [""] * (4 - len(p["alt"]))
        row = [p["key"], p["weekday"], p["key"][:7], FORMAT[p["format"]], p["mode"], p["id"], p["author"],
               p["path"], p["occasion"] + (" (revisit)" if p["revisit"] else ""), p["folder"], p["caption"],
               *alts, kept.get((p["key"], p["id"]), "")]
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FONT, size=10, bold=bool(p["occasion"]) and i <= 9)
            c.alignment = Alignment(vertical="top", wrap_text=i in (11, 12, 13, 14, 15))
            if p["mode"] == "dark" and i <= 10:
                c.fill = PatternFill("solid", fgColor=STONE)
        ws.cell(r, 16).fill = PatternFill("solid", fgColor=YELLOW)
        ws.row_dimensions[r].height = 150
    dv = DataValidation(type="list", formula1='"yes"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"P2:P{len(posts) + 1}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:P{len(posts) + 1}"
    wb.calculation.fullCalcOnLoad = True  # the count on "How to use"
    wb.save(XLSX)
    print(f"{XLSX.relative_to(OUT.parent.parent)}: {len(posts)} posts, {len(kept)} marked scheduled")


if __name__ == "__main__":
    main()
