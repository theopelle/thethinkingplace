"""One-time: turn the two-year calendar into the 728-day looping cycle (22 October 2026).

    python3 tools/make_loop_calendar.py

Reads content/calendar.xlsx as the content team made it, keeps a copy in exports/, and
writes the new content/calendar.xlsx (Calendar, Pinned dates, How it works, By month)
plus a change log for the team in exports/. After this, the team edits the spreadsheet
as before; the build checks every rule (tools/calendar_rules.py).

How the cycle is made:
- The first 62 days (to 20 December 2026) stay exactly as planned.
- The six pinned thoughts leave the cycle.
- Revisit days: every Sunday from day 189 (late April 2027), plus Wednesdays spread evenly.
- New thoughts keep the team's order: difficult ones (4–5) on Saturdays, spread evenly,
  the others on the remaining days. Small swaps then restore "no same path or author
  within 4 days" where moving things created a clash.
- Revisits: a target per difficulty level in proportion to the cycle, at least one each.
  Difficult revisits only on Sundays after a Saturday that isn't difficult. Within a
  level, the highest value first, as close to a year after the first showing as possible,
  never within 189 days of another showing of the same thought.
- Evening teasers stay with the thought they describe; new ones only where none exists.
"""
import collections, datetime as dt, os, shutil, sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from content_io import ROOT, read_thoughts
import calendar_rules as R

PINNED = [  # date, thought, occasion (as in "a thought for …"), evening teaser
    ("01-01", 579, "the New Year", "Tomorrow: a thought for the New Year."),
    ("03-20", 538, "the spring equinox", "Tomorrow: a thought for the spring equinox."),
    ("04-05", 517, "blossom time", "Tomorrow: a thought for blossom time."),
    ("06-21", 347, "midsummer", "Tomorrow: a thought for midsummer."),
    ("09-22", 31, "the autumn equinox", "Tomorrow: a thought for the autumn equinox."),
    ("12-21", 152, "the longest night", "Tomorrow: a thought for the longest night."),
]
TODAY = "2026-10-22"
SRC = os.path.join(ROOT, "content", "calendar.xlsx")


def read_old():
    ws = load_workbook(SRC, data_only=True)["Calendar"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    return [{"date": r[0].date(), "id": int(r[2]), "revisit": r[8] == "Yes", "special": r[10], "teaser": r[11]}
            for r in rows]


def build(old, T):
    pinned = {md: i for md, i, _, _ in PINNED}
    pinned_ids = set(pinned.values())
    L, F = R.CYCLE, R.LAUNCH_DAYS
    frozen = [d["id"] for d in old[:F]]
    assert not any(d["revisit"] for d in old[:F]) and not pinned_ids & set(frozen)
    later = [d["id"] for d in old if not d["revisit"] and d["id"] not in pinned_ids and d["id"] not in frozen]
    D = [i for i in later if R.hard(T[i])]
    N = [i for i in later if not R.hard(T[i])]
    wd = lambda k: R.date_of(k).weekday()

    # Revisit days: Sundays from day 189, plus evenly spread Wednesdays.
    n_rev = L - (len(frozen) + len(later))
    sundays = [k for k in range(R.REVISIT_GAP, L) if wd(k) == 6]
    weds = [k for k in range(R.REVISIT_GAP, L) if wd(k) == 2]
    n_wed = n_rev - len(sundays)
    rev = set(sundays) | {weds[round(j * (len(weds) - 1) / (n_wed - 1))] for j in range(n_wed)}
    assert len(rev) == n_rev

    # Revisit targets per difficulty level: in proportion, at least one each.
    share = collections.Counter(T[i]["difficulty"] for i in frozen + later)
    raw = {lv: n_rev * c / len(frozen + later) for lv, c in share.items()}
    target = {lv: max(1, int(x)) for lv, x in raw.items()}
    for lv in sorted(raw, key=lambda lv: raw[lv] - int(raw[lv]), reverse=True)[:n_rev - sum(target.values())]:
        target[lv] += 1
    n_hard_rev = sum(c for lv, c in target.items() if lv >= 4)

    # Difficult new thoughts on Saturdays, evenly; enough plain Saturdays in the revisit
    # season for every difficult revisit to follow one.
    # A difficult revisit needs a difficult thought shown 189 days before it: the first one
    # is in the launch days, so difficult revisits can start 189 days after it.
    first_hard = min(k for k, i in enumerate(frozen) if R.hard(T[i]))
    zone = first_hard + R.REVISIT_GAP
    early = [k for k in range(F, zone) if wd(k) == 5]
    late = [k for k in range(zone, L) if wd(k) == 5]
    n_late = min(len(late) - n_hard_rev, round(len(D) * len(late) / (len(early) + len(late))))
    n_early = len(D) - n_late
    assert n_early <= len(early)
    spread = lambda pool, n: [pool[round(j * (len(pool) - 1) / max(1, n - 1))] for j in range(n)] if n else []
    hard_sats = set(spread(early, n_early)) | set(spread(late, n_late))
    plain_late_sats = [k for k in late if k not in hard_sats]

    cycle = [(i, False) for i in frozen] + [None] * (L - F)
    dq, nq = iter(D), iter(N)
    for k in range(F, L):
        if k in rev:
            continue
        cycle[k] = (next(dq) if k in hard_sats else next(nq), False)
    assert next(dq, None) is None and next(nq, None) is None

    def clash(k, cyc):
        """Same path or author as a day within the 4-day window (empty days ignored)."""
        if not cyc[k]:
            return False
        a = T[cyc[k][0]]
        for d in range(-R.WINDOW + 1, R.WINDOW):
            o = cyc[(k + d) % L]
            if d and o and (T[o[0]]["path"] == a["path"] or R.author(T[o[0]]["source"]) == R.author(a["source"])):
                return True
        return False

    def repair(cyc, movable):
        """Swap clashing days with a nearby day of the same kind until nothing clashes."""
        count = lambda: sum(clash(k, cyc) for k in range(L))
        for _ in range(500):
            now = count()
            if not now:
                return
            for k in [k for k in range(L) if clash(k, cyc) and movable(k)]:
                best = None
                for d in sorted(range(-28, 29), key=abs):
                    j = k + d
                    if d == 0 or not 0 <= j < L or not movable(j) or (k in hard_sats) != (j in hard_sats):
                        continue
                    cyc[k], cyc[j] = cyc[j], cyc[k]
                    c = count()
                    cyc[k], cyc[j] = cyc[j], cyc[k]
                    if c < now and (best is None or c < best[0]):
                        best = (c, j)
                        if c == 0 or abs(d) > 7:
                            break
                if best:
                    j = best[1]
                    cyc[k], cyc[j] = cyc[j], cyc[k]
                    break
            else:
                raise SystemExit(f"cannot repair {now} clashes, e.g. {[R.date_of(k) for k in range(L) if clash(k, cyc)][:5]}")
        raise SystemExit("repair did not finish")

    repair(cycle, lambda k: k >= F and k not in rev and cycle[k] is not None)

    # Revisits. Difficult ones first, on Sundays after a plain Saturday.
    hard_slots = spread([k + 1 for k in plain_late_sats if k + 1 in rev], n_hard_rev)
    fits = lambda p: any(R.REVISIT_GAP <= k - p <= L - R.REVISIT_GAP for k in hard_slots)
    # A level with very few thoughts (level 5 has one) needs them placed early enough in
    # the cycle for a revisit to fit: swap with an earlier difficult Saturday if needed.
    for lv in (lv for lv in target if lv >= 4 and share[lv] <= target[lv]):
        for k in [k for k, c in enumerate(cycle) if c and T[c[0]]["difficulty"] == lv and not fits(k)]:
            for j in sorted((j for j in hard_sats if fits(j) and T[cycle[j][0]]["difficulty"] != lv), key=lambda j: abs(j - k)):
                cycle[k], cycle[j] = cycle[j], cycle[k]
                if not clash(k, cycle) and not clash(j, cycle):
                    break
                cycle[k], cycle[j] = cycle[j], cycle[k]
            else:
                raise SystemExit(f"no Saturday lets #{cycle[k][0]} be revisited")
    pos = {c[0]: k for k, c in enumerate(cycle) if c}
    used, got = set(), collections.Counter()

    def place(k, levels):
        cand = [i for i in pos if i not in used and T[i]["difficulty"] in levels
                and R.REVISIT_GAP <= k - pos[i] <= L - R.REVISIT_GAP]
        for i in sorted(cand, key=lambda i: (-T[i]["value"], abs(k - pos[i] - 364), i)):
            cycle[k] = (i, True)
            if not clash(k, cycle):
                used.add(i); got[T[i]["difficulty"]] += 1
                return True
        cycle[k] = None
        return False

    left = list(hard_slots)
    for lv in sorted((lv for lv in target if lv >= 4), key=lambda lv: target[lv]):  # rarest level first
        for k in list(left):
            if got[lv] == target[lv]:
                break
            if place(k, [lv]):
                left.remove(k)
    for k in left:
        if not place(k, [4, 5]):
            raise SystemExit(f"no difficult revisit fits {R.date_of(k)}")
    for k in sorted(rev - set(hard_slots)):
        order = sorted((lv for lv in target if lv < 4), key=lambda lv: got[lv] - target[lv])
        if not any(place(k, [lv]) for lv in order):
            raise SystemExit(f"no revisit fits {R.date_of(k)}")
    return cycle, pinned, target, got


def teasers_by_thought(old):
    """Each teaser belongs to the thought shown the next day."""
    out = {}
    for prev, cur in zip(old, old[1:]):
        if not cur["revisit"] and prev["teaser"]:
            out[cur["id"]] = prev["teaser"]
    return out


def kind_words(e, words):
    who, on = words["paths"].get(e["path"]), words["themes"].get(e["theme1"])
    return f"Tomorrow: {who} on {on}." if who and on else ""


def main():
    T = {t["id"]: t for t in read_thoughts()}
    old = read_old()
    if len(old) != 731:
        sys.exit("content/calendar.xlsx is no longer the two-year calendar: nothing to do")
    cycle, pinned, target, got = build(old, T)
    problems, stats = R.check(cycle, pinned, T)
    if problems:
        sys.exit("\n".join(problems[:30]))

    # Teasers: kept with their thought; new only where none exists.
    keep = teasers_by_thought(old)
    import json
    words = json.load(open(os.path.join(ROOT, "calendar.json"), encoding="utf-8"))["words"]
    new_teasers = []
    teaser_rows = []
    for k in range(R.CYCLE):
        nxt_id, nxt_rev = cycle[(k + 1) % R.CYCLE]
        if nxt_rev:
            text = R.REVISIT_TEASER
        elif nxt_id in keep:
            text = keep[nxt_id]
        else:
            text = kind_words(T[nxt_id], words)
            new_teasers.append((nxt_id, text))
        teaser_rows.append(text)

    backup = os.path.join(ROOT, "exports", f"calendar-before-loop-{TODAY}.xlsx")
    shutil.copy(SRC, backup)
    write_calendar(cycle, pinned, T, teaser_rows, stats, target, got)
    write_changes(old, cycle, T, new_teasers, keep, stats, target, got)
    print(f"cycle: {R.CYCLE} days, {sum(not r for _, r in cycle)} new + {sum(r for _, r in cycle)} revisits; "
          f"pinned {len(pinned)}; revisits by difficulty {dict(sorted(got.items()))} (target {dict(sorted(target.items()))})")
    print(f"shortest gap over {stats['years']} years: {stats['shortest_gap'][0]} days; {stats['kinds']}")
    print(f"old calendar kept in {os.path.relpath(backup, ROOT)}")


HEAD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="EFEAE2")
FILLS = {"launch": "E2EFDA", "hard": "FCE4D6", "revisit": "DDEBF7", "pinned": "FFF2CC"}
TOP = Alignment(vertical="top", wrap_text=True)


def sheet(ws, header, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, TOP
    for col, w in zip("ABCDEFGHIJKL", widths):
        ws.column_dimensions[col].width = w


def short(text, n=140):
    text = text.replace("\n", " / ")
    return text if len(text) <= n else text[:n - 1] + "…"


def write_calendar(cycle, pinned, T, teasers, stats, target, got):
    wb = Workbook()
    ws = wb.active
    ws.title = "Calendar"
    sheet(ws, ["Date (cycle 1)", "Weekday", "Thought ID", "Path", "Author + book", "Thought", "Main theme",
               "Difficulty", "Revisit", "First shown", "Evening teaser (shown this evening, about tomorrow)"],
          [11, 10, 9, 22, 32, 56, 22, 9, 8, 11, 44])
    pos = {i: k for k, (i, rev) in enumerate(cycle) if not rev}
    for k, (i, rev) in enumerate(cycle):
        t, d = T[i], R.date_of(k)
        ws.append([d, d.strftime("%A"), i, t["path"], t["source"], short(t["text"]), t["theme1"], t["difficulty"],
                   "Yes" if rev else None, R.date_of(pos[i]) if rev else None, teasers[k]])
        fill = "revisit" if rev else "hard" if R.hard(t) else "launch" if d < dt.date(2026, 11, 1) else None
        for c in ws[ws.max_row]:
            c.alignment = TOP
            if fill:
                c.fill = PatternFill("solid", fgColor=FILLS[fill])
        ws.cell(ws.max_row, 1).number_format = ws.cell(ws.max_row, 10).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:K{ws.max_row}"

    ws = wb.create_sheet("Pinned dates")
    sheet(ws, ["Date (day-month)", "Thought ID", "Occasion (a thought for …)", "Evening teaser (the evening before)",
               "Path", "Author + book", "Thought"], [16, 10, 24, 44, 22, 32, 56])
    for md, i, occasion, teaser in PINNED:
        t = T[i]
        ws.append([md[3:] + "-" + md[:2], i, occasion, teaser, t["path"], t["source"], short(t["text"])])
        for c in ws[ws.max_row]:
            c.alignment, c.fill = TOP, PatternFill("solid", fgColor=FILLS["pinned"])
    ws.freeze_panes = "C2"

    ws = wb.create_sheet("How it works")
    ws.column_dimensions["A"].width = 120
    k = stats["kinds"]
    for line in [
        "The Thinking Place publication calendar: a cycle that loops forever",
        "One cycle is 728 days (104 weeks), from Tuesday 20 October 2026 to Monday 16 October 2028; then it starts again "
        "(cycle 2 begins Tuesday 17 October 2028). Weekdays never move. The day changes at each reader's local midnight.",
        "The dates in the Calendar sheet are those of cycle 1. In cycle 2, add 728 days (the same weekday, two years minus "
        "two or three days later).",
        "Thought ID = the # column in thoughts.xlsx. These numbers are permanent.",
        None,
        "Rules (the build checks every one and stops if any is broken)",
        "1. Every thought appears once per cycle as new; the six pinned thoughts are outside the cycle.",
        "2. The first days, 20 October to 20 December 2026, are the launch days as first planned: easy, high-value thoughts "
        "first (green before 1 November).",
        "3. No same path or same author within any 4 days in a row, including from the end of the cycle to its start.",
        "4. Difficult thoughts (4–5) as new only on Saturdays; as a revisit only on a Sunday after a Saturday that is not "
        "difficult; never two difficult days in a row (orange).",
        "5. Revisits (blue): every Sunday from late April 2027, plus Wednesdays spread evenly; none in the first six months "
        "of the cycle. Each one revisits a thought first shown at least 189 days before, and at least 189 days before it "
        "comes round again in the next cycle. Chosen across all difficulty levels in proportion, the most valuable first, "
        "as close to a year after the first showing as possible.",
        "6. Pinned dates (sheet 'Pinned dates', yellow): the same thought on the same date every year. Pinned dates "
        "drift through the 728-day cycle, so on a pinned date the cycle's thought of that day moves to the next revisit "
        "day if there is one within a week (that revisit is skipped), otherwise it waits for the next cycle. Keep pinned "
        "dates at least 8 days apart.",
        "7. No thought ever appears within six months of another appearance of itself, in either direction, across "
        "cycles, pinned dates and moves included.",
        "8. The evening teaser in each row is shown that evening and describes the next day's thought. A teaser belongs "
        "to the thought it describes: if thoughts are moved, move their teasers with them. Before a revisit: 'Tomorrow: a "
        "thought you may have met before.'",
        None,
        "Checks (22 October 2026)",
        f"Revisits by difficulty: {dict(sorted(got.items()))} (target {dict(sorted(target.items()))}).",
        f"Checked day by day over {stats['years']} years: shortest gap between two showings of a thought "
        f"{stats['shortest_gap'][0]} days. Pinned days: {k.get('pinned on a revisit', 0)} on a revisit, "
        f"{k.get('moved', 0)} moved their thought, {k.get('skipped (waits for the next cycle)', 0)} waited for the next cycle.",
    ]:
        ws.append([line])
        ws.cell(ws.max_row, 1).alignment = TOP
        if line and (line.startswith("The Thinking Place") or line in ("Checks (22 October 2026)",) or line.startswith("Rules")):
            ws.cell(ws.max_row, 1).font = HEAD

    ws = wb.create_sheet("By month")
    paths = ["Stoicism", "Epicureanism", "Classical Greek Philosophy", "Enlightenment Ethics", "Existential Thought",
             "Humanists and Essayists", "Buddhist Philosophy", "Indian Philosophy", "Taoism", "Confucianism",
             "Japanese Aesthetics", "Persian and Sufi Poetry"]
    sheet(ws, ["Month", "New", "Revisits", "Difficult", *paths], [10, 7, 9, 9] + [12] * 8)
    months = collections.OrderedDict()
    for k, (i, rev) in enumerate(cycle):
        m = months.setdefault(R.date_of(k).strftime("%Y-%m"), collections.Counter())
        m["rev" if rev else "new"] += 1
        m["hard"] += R.hard(T[i])
        m[T[i]["path"]] += 1
    for m, c in months.items():
        ws.append([m, c["new"], c["rev"], c["hard"], *[c[p] for p in paths]])
    wb.save(SRC)


def write_changes(old, cycle, T, new_teasers, keep, stats, target, got):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 120
    old_rev = {d["id"] for d in old if d["revisit"]}
    new_rev = {i for i, r in cycle if r}
    for line in [
        "Calendar change log: the looping calendar (22 October 2026)",
        "The calendar is now a 728-day cycle that repeats forever (see 'How it works' in calendar.xlsx). The previous "
        f"calendar is kept in exports/calendar-before-loop-{TODAY}.xlsx.",
        "What changed:",
        "- 731 days became 728: the 3 days dropped were revisits (17–19 October 2028), so no thought was lost.",
        "- Six thoughts are now pinned to a date every year (sheet 'Pinned dates'): 1 January #579 Khayyam, 20 March #538 "
        "Kamo no Chōmei, 5 April #517 Kenkō (blossoms), 21 June #347 Dhammapada, 22 September #31 Boethius, 21 December "
        "#152 Kant.",
        "- The first 62 days (to 20 December 2026) are unchanged. After that, new thoughts keep your order; difficult ones "
        "stay on Saturdays; a few were swapped with a neighbour to keep 'no same path or author within 4 days'.",
        f"- Revisits: {len(new_rev)} instead of 127, every Sunday from late April 2027 plus Wednesdays, no longer a block "
        f"at the end; across all difficulty levels {dict(sorted(got.items()))}. {len(new_rev & old_rev)} thoughts were "
        f"revisited before too, {len(new_rev - old_rev)} are newly revisited, {len(old_rev - new_rev)} are no longer "
        "revisited (sheet 'Revisits').",
        "- The 'Special day' column is gone: after the first cycle the labels no longer fall on their dates. Only the "
        "pinned dates stay tied to a date (sheet 'Special days removed').",
        f"- Evening teasers stay with the thought they describe. {len(new_teasers) + 6} new teasers to review (sheet "
        "'New teasers').",
        "- Please check the thoughts that seem tied to a date or season (sheet 'Seasonal thoughts').",
    ]:
        ws.append([line])
        ws.cell(ws.max_row, 1).alignment = TOP
    ws.cell(1, 1).font = HEAD

    ws = wb.create_sheet("New teasers")
    sheet(ws, ["Thought ID", "Shown", "New teaser", "Thought"], [10, 40, 50, 80])
    for i, text in new_teasers:
        ws.append([i, "the evening before it, in the cycle", text, short(T[i]["text"])])
    for md, i, occ, teaser in PINNED:
        ws.append([i, f"the evening before {md[3:]}-{md[:2]} (pinned)", teaser, short(T[i]["text"])])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = TOP

    ws = wb.create_sheet("Revisits")
    sheet(ws, ["Date (cycle 1)", "Weekday", "Thought ID", "Difficulty", "Value", "First shown", "Days after",
               "Revisited before?", "Thought"], [11, 10, 9, 9, 7, 11, 9, 14, 80])
    pos = {i: k for k, (i, r) in enumerate(cycle) if not r}
    for k, (i, r) in enumerate(cycle):
        if r:
            ws.append([R.date_of(k), R.date_of(k).strftime("%A"), i, T[i]["difficulty"], T[i]["value"],
                       R.date_of(pos[i]), k - pos[i], "yes" if i in old_rev else "new", short(T[i]["text"], 100)])
            ws.cell(ws.max_row, 1).number_format = ws.cell(ws.max_row, 6).number_format = "yyyy-mm-dd"
    ws.append([])
    ws.append(["No longer revisited:", None, ", ".join(f"#{i}" for i in sorted(old_rev - new_rev))])

    ws = wb.create_sheet("Special days removed")
    sheet(ws, ["Old date", "Special day", "Thought ID", "Now", "Thought"], [11, 26, 10, 30, 80])
    pinned_ids = {i for _, i, _, _ in PINNED}
    for d in old:
        if d["special"]:
            ws.append([d["date"], d["special"], d["id"], "pinned" if d["id"] in pinned_ids else "in the cycle",
                       short(T[d["id"]]["text"], 100)])
            ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd"

    ws = wb.create_sheet("Seasonal thoughts")
    sheet(ws, ["Thought ID", "Why", "Thought"], [10, 50, 100])
    for i, why in [(486, "the year turns cold (winter)"), (536, "Rikiu sweeping autumn leaves"),
                   (543, "Sei Shōnagon, winter dawn (was on Women's Day in March)"), (464, "frost and snow"),
                   (517, "blossoms: now pinned to 5 April")] + \
                  [(i, "seasons only as an image: fine any time") for i in (8, 40, 76, 110, 342, 343, 429, 437, 440, 520, 546, 580, 603)]:
        ws.append([i, why, short(T[i]["text"], 160)])
    for name in wb.sheetnames:
        for row in wb[name].iter_rows(min_row=2):
            for c in row:
                c.alignment = TOP
    wb.save(os.path.join(ROOT, "exports", f"calendar-loop-changes-{TODAY}.xlsx"))


if __name__ == "__main__":
    main()
