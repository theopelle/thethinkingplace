"""The Library's pages: content/pages.xlsx (+ content/page-related.xlsx) -> pages.json, with the checks.

    python3 tools/build_pages.py        (also run at the end of build_content.py)

Design: exports/handover-library-ui-2026-10-30.md and the three handovers it builds on. Page text is
copied exactly as written in the sheet (light Markdown); the app never rewrites it. Every page is
written with its Status; the app shows only "Ready" pages (drafts only in test APKs). Checked:
  - every live thought belongs to exactly one thinker or text page, with the same path;
  - every thinker and text page has at least one thought;
  - every Ready page has a header (thinkers, texts and eras), a section and sources;
  - the names pages are linked by inside other pages' text point to one page only;
  - every "Related" entry in content/page-related.xlsx names a page.
"""
import collections, json, os, re, sys, unicodedata

from content_io import ROOT, CONTENT
from openpyxl import load_workbook, Workbook
import calendar_rules

# Thought sources that don't name their page ("Analects 1.4" is Confucius; the four Upanishads
# share one page). Everything else matches a page's title, its search aliases or its header name.
AUTHOR_PAGE = {"Analects": "thinker-confucius"}
UPANISHAD = re.compile(r"\bUpanishad$")
# Surnames that mostly mean someone else in the pages ("James Legge", "King James").
NOT_A_LINK = {"James"}
FIRST_SECTION, LAST_SECTION = "In one line", "Essay"
PLACEHOLDER = re.compile(r"^\*\(Built by the app:.*?\)\*\s*$", re.M)
RELATED = os.path.join(CONTENT, "page-related.xlsx")


def fold(s):
    """Accents and case ignored: 'Kenkō' and 'kenko' match."""
    return unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode().lower().strip()


def split_list(s):
    return [x.strip() for x in re.split(r"[,;\n]", s or "") if x.strip()]


def start_year(dates):
    """When an era begins, for the chronological list: "6th–3rd centuries BCE" -> -600,
    "1st century BCE–6th century CE" -> -100, "about 1600–1700" -> 1600, "late 19th–…" -> 1870."""
    m = re.search(r"(\d+)(st|nd|rd|th)?", dates)
    n, century = int(m.group(1)), bool(m.group(2))
    bce = re.search(r"\b(BCE|CE)\b", dates[m.end():])
    bce = bool(bce and bce.group(1) == "BCE")
    year = n * 100 if century and bce else (n - 1) * 100 if century else n
    year = -year if bce else year
    return year + (70 if century and re.match(r"\s*late\b", dates[:m.start()]) else 0)


def read_sheet(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if r and r[0]]


def make_related_sheet(pages):
    """content/page-related.xlsx: one row per path, era and topic page, for the content team to fill."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Related"
    ws.append(["Page ID", "Type", "Title", "Related (page titles or IDs, comma-separated, at most 6)"])
    for p in pages:
        if p["type"] in ("path", "era", "topic"):
            ws.append([p["id"], p["type"], p["title"], None])
    for col, width in zip("ABCD", (34, 10, 42, 70)):
        ws.column_dimensions[col].width = width
    wb.save(RELATED)
    print(f"created {os.path.relpath(RELATED, ROOT)} (empty: the content team fills it)")


def main():
    wb = load_workbook(os.path.join(CONTENT, "pages.xlsx"), data_only=True)
    head = [c.value for c in next(wb["Pages"].iter_rows(max_row=1))]
    section_cols = head[head.index(FIRST_SECTION):head.index(LAST_SECTION) + 1]
    sources = collections.defaultdict(list)
    for s in read_sheet(wb, "Sources"):  # a source without a link ("The app's thoughts #4, …") shows as text
        sources[s["Page ID"]].append([s["Source"], s.get("Link") or ""])
    problems, pages = [], []

    for r in read_sheet(wb, "Pages"):
        pid, kind = r["Page ID"].strip(), r["Type"].strip()
        # Header: "**Seneca** · Córdoba, Spain · about 4 BCE – 65 CE · Stoicism"; the name is the title.
        parts = [x.strip() for x in (r["Header"] or "").split("·")]
        name = parts[0].strip("* ") if parts and parts[0].startswith("*") else None
        header = parts[1:] if name else [x for x in parts if x]
        sections = []
        for col in section_cols:
            md = PLACEHOLDER.sub("", r.get(col) or "").strip()
            if md:
                sections.append([col, md])
        pages.append({"id": pid, "type": kind, "title": r["Title"].strip(), "status": (r["Status"] or "").strip(),
                      "header": header if kind != "topic" else [], "path": r["Path"] or "", "era": r["Era"] or "",
                      "eraOrder": r["Era order"], "theme": r["Theme"] or "", "society": (r["Society topic"] or "") == "yes",
                      "thoughts": [int(x) for x in split_list(str(r["Thought IDs"] or ""))],
                      "tags": split_list(r["Tags"]), "aliases": split_list(r["Search aliases"]),
                      "name": name, "sections": sections, "sources": sources.get(pid, [])})
    by_id = {p["id"]: p for p in pages}

    # Thoughts -> thinker and text pages.
    entries = [e for e in json.load(open(os.path.join(ROOT, "entries.json"), encoding="utf-8")) if not e.get("retired")]
    names = {}
    for p in pages:
        if p["type"] in ("thinker", "text"):
            for n in [p["title"], re.sub(r"^The ", "", p["title"]), p["name"] or "", *p["aliases"]]:
                if n:
                    names[fold(n.strip("*"))] = p["id"]
    for e in entries:
        a = calendar_rules.author(e["source"])
        pid = AUTHOR_PAGE.get(a) or ("text-the-upanishads" if UPANISHAD.search(a) else names.get(fold(a)))
        if not pid:
            problems.append(f"#{e['id']} ({e['source']}): no thinker or text page")
            continue
        if by_id[pid]["path"] != e["path"]:
            problems.append(f"#{e['id']} ({e['source']}): path {e['path']!r}, but its page {pid} says {by_id[pid]['path']!r}")
        by_id[pid]["thoughts"].append(e["id"])
    for p in pages:
        if p["type"] in ("thinker", "text") and not p["thoughts"]:
            problems.append(f"{p['id']}: no thought")
        if p["type"] == "path":  # a path page's thoughts: all of that path's
            p["thoughts"] = [e["id"] for e in entries if e["path"] == p["title"]]
            if not p["thoughts"]:
                problems.append(f"{p['id']}: no thought has this path")
        if p["type"] == "era":
            p["start"] = start_year(p["header"][0] if p["header"] else "")
        if p["type"] in ("thinker", "text"):  # for "The thinkers through time": birth, or a text's start
            p["start"] = start_year(p["header"][1] if len(p["header"]) > 1 else "")
    # Topic pages: their hand-picked thoughts, plus every thought of their theme (Freedom has both).
    live = {e["id"]: e for e in entries}
    themes = collections.Counter(t for e in entries for t in e.get("themes", []))
    for p in pages:
        if p["type"] != "topic":
            continue
        for i in p["thoughts"]:
            if i not in live:
                problems.append(f"{p['id']}: thought #{i} is not a live thought")
        if p["theme"] and p["theme"] not in themes:
            problems.append(f"{p['id']}: theme {p['theme']!r} is not a theme of any thought")
        p["thoughts"] = sorted(set(p["thoughts"]) | {e["id"] for e in entries if p["theme"] in e.get("themes", [])})
        if not p["thoughts"]:
            problems.append(f"{p['id']}: no thought")
    for t in themes:
        n = sum(p["type"] == "topic" and p["theme"] == t for p in pages)
        if n != 1:
            problems.append(f"theme {t!r} has {n} pages (exactly one expected)")

    # Ready pages are complete.
    for p in pages:
        if p["status"] != "Ready":
            continue
        if not p["header"] and p["type"] not in ("essay", "topic", "path"):  # paths and topics have no header line
            problems.append(f"{p['id']} is Ready but has no header")
        if not p["sections"]:
            problems.append(f"{p['id']} is Ready but has no section")
        if not p["sources"] and p["type"] != "essay":
            problems.append(f"{p['id']} is Ready but has no sources")

    # Themes -> their page (a theme page, or a society page that doubles as one).
    theme_page = {}
    for p in pages:
        if p["type"] == "topic" and p["theme"]:
            theme_page[p["theme"]] = p["id"]
    # Related, for thinkers and texts: the theme pages that come up most in their thoughts.
    thoughts = {e["id"]: e for e in entries}
    related = {}
    for p in pages:
        if p["type"] in ("thinker", "text"):
            count = collections.Counter(t for i in p["thoughts"] for t in thoughts[i].get("themes", []) if t in theme_page)
            related[p["id"]] = [theme_page[t] for t, _ in count.most_common(6)]
    # ...and for the other pages, what the content team writes in content/page-related.xlsx.
    if not os.path.exists(RELATED):
        make_related_sheet(pages)
    titles = {fold(p["title"]): p["id"] for p in pages}
    for row in read_sheet(load_workbook(RELATED, data_only=True), "Related"):
        got = []
        for x in split_list(row.get("Related (page titles or IDs, comma-separated, at most 6)")):
            pid = x if x in by_id else titles.get(fold(x))
            if not pid:
                problems.append(f"page-related.xlsx, {row['Page ID']}: {x!r} is not a page")
            elif pid != row["Page ID"]:
                got.append(pid)
        if got:
            related[row["Page ID"]] = got[:6]

    # Names that link to a page from inside other pages' text: thinkers, texts and paths only
    # (never eras or themes), by title, search aliases, and surname when it is unambiguous.
    link = collections.defaultdict(set)
    for p in pages:
        if p["type"] in ("thinker", "text", "path"):
            for n in [p["title"], re.sub(r"^The ", "", p["title"]), *p["aliases"]]:
                link[n].add(p["id"])
    surnames = collections.defaultdict(set)
    for p in pages:
        if p["type"] == "thinker" and " " in p["title"]:
            surnames[p["title"].split()[-1]].add(p["id"])
    for n, ids in surnames.items():
        if len(ids) == 1 and n not in link and len(n) > 3 and n not in NOT_A_LINK:
            link[n] |= ids
    links = {n: next(iter(ids)) for n, ids in link.items() if len(ids) == 1}

    # Glossary cards -> pages: author and book cards open their thinker or text page.
    cards = {}
    for g in json.load(open(os.path.join(ROOT, "glossary.json"), encoding="utf-8")):
        if g.get("school") or g.get("concept"):
            continue
        for n in [g["term"], *g.get("match", [])]:
            pid = AUTHOR_PAGE.get(n) or ("text-the-upanishads" if UPANISHAD.search(n) else names.get(fold(n)))
            if pid:
                cards[g["term"]] = pid
                break
    glossary_links = {g["Concept card"]: g["Page ID"] for g in read_sheet(wb, "Glossary links") if g["Page ID"] in by_id}

    if problems:
        sys.exit("library pages:\n" + "\n".join(problems))
    for p in pages:
        p.pop("name")
        if not p["thoughts"]:
            p.pop("thoughts")
    out = {"pages": pages, "related": related, "links": links, "cards": cards, "glossaryLinks": glossary_links}
    json.dump(out, open(os.path.join(ROOT, "pages.json"), "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    kinds = collections.Counter(p["type"] for p in pages)
    print(f"{len(pages)} library pages ({sum(p['status'] == 'Ready' for p in pages)} ready: "
          f"{', '.join(f'{n} {k}' for k, n in kinds.most_common())}), every thought on its page -> pages.json")


if __name__ == "__main__":
    main()
