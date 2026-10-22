"""Read the content team's spreadsheets in content/ (the master copy of all text)."""
import os
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTENT = os.path.join(ROOT, "content")

COLS = ["id", "text", "source", "context", "translation_raw", "prompt", "path", "status",
        "theme1", "theme2", "difficulty", "value", "to_check"]

# Better sources for links that point at a blog or a page without the full text.
# The quote is checked against (and credited to) the replacement.
M = "https://neonvagabond.xyz/mirrors/sacred-texts/"
EXTRA_SOURCES = {
    # sacred-texts.com refuses scripts; a full mirror of it lives on neonvagabond.xyz.
    # "crawl:" fetches every page linked from an index and joins them.
    **{f"https://sacred-texts.com/isl/masnavi/msn0{i}.htm": M + f"isl/masnavi/msn0{i}.htm" for i in range(1, 7)},
    "https://sacred-texts.com/isl/taa/taa14.htm": M + "isl/taa/taa14.htm",
    **{u: "crawl:" + M + "hin/sbe15/index.htm" for u in [
        "https://sacred-texts.com/book/the-upanishads-part-ii/i-2",
        "https://sacred-texts.com/book/the-upanishads-part-ii/i-3",
        "https://sacred-texts.com/hin/sbe15/sbe15076.htm",
        "https://sacred-texts.com/hin/sbe15/sbe15020.htm"]},
    "https://hinduwebsite.com/sacredscripts/hinduism/upanishads/vagasaneyi.asp": "crawl:" + M + "hin/sbe01/index.htm",
    **{f"https://sacred-texts.com/book/the-bhagavadgita/{c}": "crawl:" + M + "hin/sbe08/index.htm"
       for c in ["chapter-ii", "chatper-vi", "chapter-iii", "chapter-xii"]},
    M + "tao/sbe39/index.htm": "crawl:" + M + "tao/sbe39/index.htm",
    M + "tao/sbe40/index.htm": "crawl:" + M + "tao/sbe40/index.htm",
    "https://standardebooks.org/ebooks/leo-tolstoy/a-confession/aylmer-maude":
        "https://standardebooks.org/ebooks/leo-tolstoy/a-confession/aylmer-maude/text/single-page",
    # Secondary pages (blogs quoting the book, reading sites): check against the book itself.
    "https://rohanmaitzen.com/2009/05/06/the-other-sides-of-silence/": "https://www.gutenberg.org/ebooks/145",
    "https://georgeeliotarchive.org/items/show/58607": "https://www.gutenberg.org/ebooks/145",
    "https://apposition.substack.com/p/unvisited-tombs": "https://www.gutenberg.org/ebooks/145",
    # ctext.org blocks scripts and hinduwebsite.com refuses them: Legge's Mencius from archive.org.
    "https://ctext.org/mengzi/jin-xin-i/ens": "https://archive.org/details/fourbooksconfuci00leggiala",
    "https://www.hinduwebsite.com/sacredscripts/mencius.asp": "https://archive.org/details/fourbooksconfuci00leggiala",
}


def read_thoughts():
    ws = load_workbook(os.path.join(CONTENT, "thoughts.xlsx"), data_only=True)["Thoughts"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        t = dict(zip(COLS, r))
        t["id"] = int(t["id"])
        rows.append(t)
    return rows


def translation(t):
    raw = (t["translation_raw"] or "").strip()
    return raw.split("|")[0].strip() if raw else ""


def source_url(t):
    raw = t["translation_raw"] or ""
    url = raw.split("|")[-1].strip() if "|" in raw else ""
    return EXTRA_SOURCES.get(url, url) if url else ""


def read_calendar():
    """The 728-day cycle (dates are those of cycle 1), read by column name."""
    ws = load_workbook(os.path.join(CONTENT, "calendar.xlsx"), data_only=True)["Calendar"]
    head = [c.value for c in ws[1]]
    col = lambda name: next(i for i, h in enumerate(head) if h and h.startswith(name))
    d, i, rv, ts = col("Date"), col("Thought ID"), col("Revisit"), col("Evening teaser")
    days = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        days.append({"date": r[d].strftime("%Y-%m-%d"), "id": int(r[i]), "revisit": r[rv] == "Yes",
                     "teaser": r[ts]})
    return days


def read_pinned():
    """Pinned dates: {"MM-DD": {"id", "occasion", "teaser"}} (the sheet writes day-month)."""
    ws = load_workbook(os.path.join(CONTENT, "calendar.xlsx"), data_only=True)["Pinned dates"]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        day, month = str(r[0]).strip().split("-")
        out[f"{int(month):02d}-{int(day):02d}"] = {"id": int(r[1]), "occasion": r[2], "teaser": r[3]}
    return out


def read_glossary():
    wb = load_workbook(os.path.join(CONTENT, "glossary.xlsx"), data_only=True)
    out = {}
    for name in ["Schools", "Core concepts", "Other terms", "Map cards"]:
        ws = wb[name]
        head = [c.value for c in ws[1]]
        out[name] = [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    return out
